"""Write the review as a spreadsheet in the team's ESG DDQ checklist layout."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import Review

COLUMNS = [
    ("ESG Consideration", 48),
    ("Response", 10),
    ("Existing Conditions", 60),
    ("Mitigation / Handling", 50),
    ("Priority", 11),
    ("Type", 16),
    ("Risk", 11),
    ("Cost Estimate", 28),
    ("Sources", 30),
    ("Pre/Post Site Visit (Hide for report)", 14),
    ("Verifier", 40),
]
SECTION_FILL = PatternFill("solid", fgColor="D9E1D2")
HEADER_FILL = PatternFill("solid", fgColor="EFEFEF")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


def write_report(review: Review, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "ESG DDQ"
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    section = None
    for row in review.rows:
        if row.item.section != section:
            section = row.item.section
            if ws.max_row > 1:
                ws.append([])
            ws.append([section])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=ws.max_row, column=c).fill = SECTION_FILL
            ws.append([name for name, _ in COLUMNS])
            for c in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=ws.max_row, column=c)
                cell.font, cell.fill = Font(bold=True), HEADER_FILL
        a = row.answer
        conditions = a.existing_conditions + ("\n[Confirm on site visit]" if a.confirm_on_site else "")
        verifier = "Verified" if row.verified else f"NEEDS REVIEW: {row.verifier_note}"
        ws.append([
            row.item.question, a.response, conditions, a.mitigation, a.priority, row.item.type,
            a.risk, a.cost_estimate, "; ".join(a.sources), row.item.timing, verifier,
        ])
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=ws.max_row, column=c).alignment = WRAP
        if not row.verified:
            ws.cell(row=ws.max_row, column=len(COLUMNS)).fill = REVIEW_FILL

    _summary_sheet(wb.create_sheet("Summary", 0), review)

    data = wb.create_sheet("Data sources")
    data.append(["Connector", "Status", "Summary", "Source"])
    for c in review.connectors:
        data.append([c.connector, c.status.value, c.summary, c.source])
    data.append([])
    data.append(["Research notes"])
    for line in (review.research_notes or "").splitlines():
        data.append([line[:32000]])
    for col, width in zip("ABCD", (16, 16, 90, 40)):
        data.column_dimensions[col].width = width
    for r in data.iter_rows():
        for cell in r:
            cell.alignment = WRAP

    raw = wb.create_sheet("Connector data")
    raw.append(["Connector", "Field", "Value"])
    for c in review.connectors:
        for key, value in _flatten(c.data):
            raw.append([c.connector, key, value])
    for col, width in zip("ABC", (16, 40, 90)):
        raw.column_dimensions[col].width = width

    prop = wb.create_sheet("Property")
    for key, value in review.profile.model_dump(exclude_none=True).items():
        if value in ([], ""):
            continue
        prop.append([key, "; ".join(map(str, value)) if isinstance(value, list) else value])
    prop.column_dimensions["A"].width, prop.column_dimensions["B"].width = 26, 100

    wb.active = 0
    path = Path(path)
    wb.save(path)
    return path


def _flatten(value, prefix: str = ""):
    if isinstance(value, dict):
        for k, v in value.items():
            yield from _flatten(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(value, list) and any(isinstance(v, (dict, list)) for v in value):
        for i, v in enumerate(value):
            yield from _flatten(v, f"{prefix}[{i}]")
    else:
        text = "; ".join(map(str, value)) if isinstance(value, list) else value
        yield prefix, text if isinstance(text, (int, float, bool)) or text is None else str(text)[:32000]


def _summary_sheet(ws, review: Review) -> None:
    from .store import level, summarize

    p, s = review.profile, summarize(review)
    address = ", ".join(x for x in (p.street_address, p.city, p.state, p.postal_code) if x)
    ws.append([p.name or address or "Property"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([address])
    ws.append([])
    for label, key in [
        ("Questions answered", "questions"), ("Verified against sources", "verified"),
        ("Need human review", "needs_review"), ("High priority or risk", "high_risk"),
        ("Moderate priority or risk", "moderate_risk"), ("Confirm on site visit", "confirm_on_site"),
        ("Items with a cost estimate", "with_cost"),
    ]:
        ws.append([label, s[key]])
    ws.append(["Data sources returning data", f"{s['connectors_ok']} of {s['connectors_total']}"])
    ws.append([])
    ws.append(["Key findings (moderate or high)"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Question", "Priority", "Risk", "Existing conditions", "Mitigation", "Cost estimate"])
    for c in range(1, 7):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    for row in review.rows:
        a = row.answer
        if level(a) != "LOW":
            ws.append([row.item.question, a.priority, a.risk, a.existing_conditions, a.mitigation, a.cost_estimate])
            for c in range(1, 7):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP
    for col, width in zip("ABCDEF", (44, 12, 12, 60, 50, 28)):
        ws.column_dimensions[col].width = width
