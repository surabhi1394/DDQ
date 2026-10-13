"""The ESG DDQ checklist template: loading it, and rebuilding it from a team spreadsheet."""

from __future__ import annotations

import json
import re
from importlib import resources
from pathlib import Path

from .models import ChecklistItem

HEADER = "ESG Consideration"
TEMPLATE_FILE = "checklist_template.json"


def load_template(path: Path | None = None) -> list[ChecklistItem]:
    if path is None:
        text = resources.files("ddq_agent.data").joinpath(TEMPLATE_FILE).read_text()
    else:
        text = Path(path).read_text()
    return [ChecklistItem(**row) for row in json.loads(text)]


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def _timing(value) -> str:
    v = _clean(value).replace(" ", "").lower()
    return {"pre": "Pre", "post": "Post", "pre/post": "Pre/Post"}.get(v, "")


def import_from_xlsx(path: Path) -> list[ChecklistItem]:
    """Read questions (not answers) from a checklist in the team's layout.

    Layout: a section title row, then a header row starting with "ESG Consideration",
    then one question per row. Column A is the question, F the type, J Pre/Post.
    A header row with no title above it continues the previous section.
    """
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [_clean(c.value) for c in next(r for r in ws.iter_rows() if _clean(r[0].value) == HEADER)]
    col = {name: i for i, name in enumerate(header) if name}
    timing_col = next((i for n, i in col.items() if n.lower().startswith("pre/post")), None)

    items: list[ChecklistItem] = []
    section = ""
    for row in ws.iter_rows(values_only=True):
        cells = [_clean(c) for c in row]
        first = cells[0] if cells else ""
        filled = [c for c in cells if c]
        if not filled or first == HEADER:
            continue
        if len(filled) == 1 and first and first.isupper():
            section = first
            continue
        if len(filled) == 1:
            continue  # stray labels such as a lone "Post"
        items.append(
            ChecklistItem(
                id=f"Q{len(items) + 1:02d}",
                section=section,
                question=first,
                type=cells[col["Type"]] if "Type" in col else "",
                timing=_timing(row[timing_col]) if timing_col is not None else "",
            )
        )
    return items


def save_template(items: list[ChecklistItem], path: Path) -> None:
    path.write_text(json.dumps([i.model_dump() for i in items], indent=2, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    import sys

    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    items = import_from_xlsx(src)
    save_template(items, dst)
    print(f"Wrote {len(items)} questions to {dst}")
