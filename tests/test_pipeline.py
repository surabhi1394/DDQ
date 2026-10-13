from types import SimpleNamespace

import openpyxl
import pytest

from ddq_agent.agents.stages import AnswerSet, RetrofitSet, RetrofitUpdate, VerdictSet
from ddq_agent.checklist import load_template
from ddq_agent.evaluate import score
from ddq_agent.llm import LLM, RefusalError
from ddq_agent.models import ChecklistAnswer, ConnectorResult, ConnectorStatus, PropertyProfile, Verdict
from ddq_agent.pipeline import run_review
from ddq_agent.report import write_report


class FakeLLM(LLM):
    """Answers every question Y; the Verifier rejects Q02 once, then accepts it."""

    def __init__(self):
        self.calls = []
        self.rejected = set()

    def text(self, system, content, tools=None):
        self.calls.append(("research", tools))
        return "Transit: TRAX station 0.3 miles away (https://example.org/trax)."

    def structured(self, system, content, output_model):
        prompt = content[-1]["text"]
        ids = [line.split("]")[0].lstrip("[") for line in prompt.splitlines() if line.startswith("[Q")]
        self.calls.append((output_model.__name__, len(ids)))
        if output_model is PropertyProfile:
            return PropertyProfile(name="60 Park", city="Salt Lake City", state="UT", year_built=2020)
        if output_model is AnswerSet:
            revised = "Reviewer feedback" in prompt
            return AnswerSet(answers=[
                ChecklistAnswer(id=i, response="Y", existing_conditions=f"Stated in OM{' (fixed)' if revised else ''}.", sources=["OM p.33"])
                for i in ids
            ])
        if output_model is RetrofitSet:
            import json
            drafts = json.loads(prompt.split("Drafted answers:\n", 1)[1])
            return RetrofitSet(updates=[
                RetrofitUpdate(id=d["id"], mitigation="No action necessary.", cost_estimate="", priority="LOW", risk="LOW", sources=[])
                for d in drafts
            ])
        if output_model is VerdictSet:
            import json
            drafts = json.loads(prompt.split("Answers to verify:\n", 1)[1])
            verdicts = []
            for d in drafts:
                reject = d["id"] == "Q02" and "Q02" not in self.rejected
                if reject:
                    self.rejected.add("Q02")
                verdicts.append(Verdict(id=d["id"], passed=not reject, issue="Unsupported claim." if reject else ""))
            return VerdictSet(verdicts=verdicts)
        raise AssertionError(output_model)


def fake_connectors(profile):
    return [ConnectorResult(connector="fines_regs", status=ConnectorStatus.OK, summary="Elevate Buildings applies.")]


@pytest.fixture
def om(tmp_path):
    path = tmp_path / "om.pdf"
    path.write_bytes(b"%PDF-1.4 fake")
    return path


def test_review_runs_workflow_and_revises_rejected_findings(om):
    llm = FakeLLM()
    items = load_template()
    review = run_review(om, llm=llm, items=items, connectors_fn=fake_connectors)

    assert len(review.rows) == len(items) == 65
    assert all(r.verified for r in review.rows)
    assert review.rounds == 1
    q02 = next(r for r in review.rows if r.item.id == "Q02")
    assert q02.answer.existing_conditions.endswith("(fixed).")
    assert q02.answer.mitigation == "No action necessary."
    # Profile, research (with web search), answers, retrofit, verify, then one revision of a single item.
    assert [c[0] for c in llm.calls] == [
        "PropertyProfile", "research", "AnswerSet", "RetrofitSet", "VerdictSet", "AnswerSet", "RetrofitSet", "VerdictSet",
    ]
    assert llm.calls[1][1][0]["type"].startswith("web_search")
    assert llm.calls[5] == ("AnswerSet", 1)


def test_unresolved_findings_are_flagged_after_max_rounds(om):
    class Stubborn(FakeLLM):
        def structured(self, system, content, output_model):
            result = super().structured(system, content, output_model)
            if output_model is VerdictSet:
                for v in result.verdicts:
                    if v.id == "Q05":
                        v.passed, v.issue = False, "Still unsupported."
            return result

    review = run_review(om, llm=Stubborn(), connectors_fn=fake_connectors, max_rounds=2)
    flagged = [r for r in review.rows if not r.verified]
    assert [r.item.id for r in flagged] == ["Q05"]
    assert review.rounds == 2


def test_report_round_trips_through_eval(om, tmp_path):
    review = run_review(om, llm=FakeLLM(), connectors_fn=fake_connectors)
    out = write_report(review, tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(out)
    ws = wb["ESG DDQ"]
    assert ws["A1"].value == "ENVIRONMENTAL & GHG PERFORMANCE"
    assert ws["A2"].value == "ESG Consideration" and ws["J2"].value.startswith("Pre/Post")
    assert ws["B3"].value == "Y"
    assert set(wb.sheetnames) == {"ESG DDQ", "Data sources", "Property"}

    # Scoring the report against itself is a perfect match.
    s = score(out, out)
    assert s.compared == 65 and s.response_accuracy == 1.0


class FakeStream:
    def __init__(self, message):
        self.message = message

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.message


class FakeClient:
    def __init__(self, messages):
        self.queue = list(messages)
        self.requests = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self._stream))

    def _stream(self, **kwargs):
        self.requests.append(kwargs)
        return FakeStream(self.queue.pop(0))


def _msg(stop_reason, text="", stop_details=None):
    return SimpleNamespace(stop_reason=stop_reason, stop_details=stop_details, content=[SimpleNamespace(type="text", text=text)])


def test_llm_resumes_paused_turns_and_sets_defaults():
    client = FakeClient([_msg("pause_turn", "searching"), _msg("end_turn", "notes")])
    assert LLM(client=client).text("sys", [{"type": "text", "text": "hi"}], tools=[{"type": "web_search_20260209"}]) == "notes"
    first, second = client.requests
    assert first["model"] == "claude-opus-5-5" and first["fallbacks"] == "default"
    assert first["thinking"] == {"type": "adaptive"}
    assert second["messages"][-1]["role"] == "assistant"


def test_llm_raises_on_refusal():
    client = FakeClient([_msg("refusal", stop_details=SimpleNamespace(explanation="declined"))])
    with pytest.raises(RefusalError):
        LLM(client=client).text("sys", [])
